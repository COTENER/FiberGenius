import io

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .models import (
    AuditoriaPuertoODF,
    DetallePuertoODF,
    FibraTramo,
    HubSite,
    InventarioFibra,
    InventarioODF,
    InventarioTramo,
    LoteImportacion,
    NodoRed,
    RackFisico,
    Ruta,
    SalaTecnica,
    TerminacionFibra,
)
from .services.puertos import conectar_puerto, reservar_puerto


ENCABEZADOS = [
    'Site', 'ODF', 'Puerto', 'Acción', 'Troncal', 'Fibra', 'Extremo',
    'Sincronizar fibra', 'Permitir mover',
]


class ImportacionOperacionesPuertosTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_superuser(
            username='admin-importacion-operaciones',
            email='operaciones@example.com',
            password='clave-segura',
        )
        cls.usuario_limitado = get_user_model().objects.create_user(
            username='limitado-importacion-operaciones',
            password='clave-segura',
        )
        cls.usuario_limitado.user_permissions.add(
            Permission.objects.get(codename='change_detallepuertoodf')
        )
        cls.site_a, cls.odf_a = cls._crear_odf('SITE-EXCEL-A', 'ODF-EXCEL-A')
        cls.site_b, cls.odf_b = cls._crear_odf('SITE-EXCEL-B', 'ODF-EXCEL-B')
        cls.puertos = [
            DetallePuertoODF.objects.create(
                odf_obj=cls.odf_a,
                puerto_odf=str(indice),
                estado_puerto='LIBRE',
            )
            for indice in range(1, 6)
        ]
        ruta = Ruta.objects.create(nombre='RUTA-EXCEL')
        nodo_a = NodoRed.objects.create(
            tipo='SITE', codigo=cls.site_a.nombre, hub_site_obj=cls.site_a
        )
        nodo_medio_1 = NodoRed.objects.create(tipo='PUNTO', codigo='EXCEL-M1')
        nodo_medio_2 = NodoRed.objects.create(tipo='PUNTO', codigo='EXCEL-M2')
        nodo_b = NodoRed.objects.create(
            tipo='SITE', codigo=cls.site_b.nombre, hub_site_obj=cls.site_b
        )
        nodos = (nodo_a, nodo_medio_1, nodo_medio_2, nodo_b)
        cls.tramos = []
        for indice in range(3):
            cls.tramos.append(InventarioTramo.objects.create(
                ruta=ruta,
                tramo_secuencia=indice + 1,
                codigo_tramo=f'RUTA-EXCEL-T{indice + 1:03d}',
                origen_nodo=nodos[indice],
                destino_nodo=nodos[indice + 1],
                capacidad_hilos=24,
            ))
        cls.fibra = InventarioFibra.objects.create(
            ruta=ruta,
            fibra_numero='F12',
            estado='DISPONIBLE',
            origen_estado='INFORMADO',
        )
        for tramo in cls.tramos:
            FibraTramo.objects.create(
                tramo=tramo,
                fibra=cls.fibra,
                numero_hilo='F12',
                estado='DISPONIBLE',
            )

    @classmethod
    def _crear_odf(cls, site_nombre, odf_nombre):
        site = HubSite.objects.create(nombre=site_nombre)
        sala = SalaTecnica.objects.create(hub_site=site, nombre=f'SALA-{odf_nombre}')
        rack = RackFisico.objects.create(sala=sala, nombre=f'RACK-{odf_nombre}')
        odf = InventarioODF.objects.create(
            rack_obj=rack,
            odf=odf_nombre,
            capacidad_puertos=24,
        )
        return site, odf

    def setUp(self):
        self.client.force_login(self.admin)

    def _excel(self, filas):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Operaciones'
        sheet.append(ENCABEZADOS)
        for fila in filas:
            sheet.append(fila)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _post(self, filas, modo='validar'):
        archivo = SimpleUploadedFile(
            'operaciones.xlsx',
            self._excel(filas),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        return self.client.post(
            reverse('api_importar_operaciones_puertos'),
            {'modo': modo, 'archivo': archivo},
        )

    def _fila(self, puerto, accion, troncal='', fibra='', extremo='', sincronizar='No', mover='No'):
        return [
            self.site_a.nombre,
            self.odf_a.odf,
            puerto,
            accion,
            troncal,
            fibra,
            extremo,
            sincronizar,
            mover,
        ]

    def test_plantilla_tiene_hojas_columnas_y_validaciones(self):
        respuesta = self.client.get(reverse('plantilla_operaciones_puertos'))
        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(respuesta.content.startswith(b'PK'))
        workbook = load_workbook(io.BytesIO(respuesta.content))
        self.assertEqual(
            workbook.sheetnames,
            ['Operaciones', 'Instrucciones', 'Ejemplo - no importar'],
        )
        self.assertEqual(
            [celda.value for celda in workbook['Operaciones'][1]],
            ENCABEZADOS,
        )
        self.assertEqual(len(workbook['Operaciones'].data_validations.dataValidation), 3)
        self.assertEqual(workbook['Operaciones'].freeze_panes, 'A2')

    def test_validacion_revierte_y_aplicacion_confirma_todo_el_archivo(self):
        filas = [
            self._fila('1', 'Reservar'),
            self._fila('2', 'Conectar', 'RUTA-EXCEL', 'F12', 'A', 'Sí'),
        ]
        validacion = self._post(filas, 'validar')
        self.assertEqual(validacion.status_code, 200)
        self.assertTrue(validacion.json()['can_apply'])
        self.puertos[0].refresh_from_db()
        self.puertos[1].refresh_from_db()
        self.fibra.refresh_from_db()
        self.assertEqual(self.puertos[0].estado_puerto, 'LIBRE')
        self.assertEqual(self.puertos[1].estado_puerto, 'LIBRE')
        self.assertEqual(self.fibra.estado, 'DISPONIBLE')
        self.assertFalse(TerminacionFibra.objects.exists())
        self.assertFalse(AuditoriaPuertoODF.objects.exists())

        aplicacion = self._post(filas, 'aplicar')
        self.assertEqual(aplicacion.status_code, 200)
        payload = aplicacion.json()
        self.assertEqual(payload['status'], 'success')
        self.assertEqual(payload['summary']['aplicadas'], 2)
        self.puertos[0].refresh_from_db()
        self.puertos[1].refresh_from_db()
        self.fibra.refresh_from_db()
        self.assertEqual(self.puertos[0].estado_puerto, 'RESERVADO')
        self.assertEqual(self.puertos[1].estado_puerto, 'OCUPADO')
        self.assertEqual(self.fibra.estado, 'OCUPADO')
        self.assertEqual(self.fibra.origen_estado, 'INFORMADO')
        self.assertEqual(
            set(self.fibra.asignaciones_tramo.values_list('estado', flat=True)),
            {'DISPONIBLE'},
        )
        self.assertTrue(
            TerminacionFibra.objects.filter(
                fibra=self.fibra,
                extremo='A',
                puerto_odf=self.puertos[1],
            ).exists()
        )
        lote = LoteImportacion.objects.get(codigo=payload['lote'])
        self.assertEqual(lote.estado, 'COMPLETADO')
        self.assertEqual(lote.filas_actualizadas, 2)
        auditorias = AuditoriaPuertoODF.objects.order_by('pk')
        self.assertEqual(auditorias.count(), 2)
        self.assertEqual(set(auditorias.values_list('origen', flat=True)), {'EXCEL'})
        self.assertEqual(
            set(auditorias.values_list('lote_importacion_id', flat=True)),
            {lote.pk},
        )

        reporte = self.client.get(payload['report_url'])
        self.assertEqual(reporte.status_code, 200)
        self.assertTrue(reporte.content.startswith(b'PK'))

    def test_error_en_una_fila_impide_aplicar_las_filas_validas(self):
        filas = [
            self._fila('1', 'Reservar'),
            self._fila('2', 'Conectar', 'RUTA-EXCEL', 'F999', 'A'),
        ]
        respuesta = self._post(filas, 'aplicar')
        self.assertEqual(respuesta.status_code, 409)
        self.assertEqual(respuesta.json()['status'], 'invalid')
        self.puertos[0].refresh_from_db()
        self.assertEqual(self.puertos[0].estado_puerto, 'LIBRE')
        self.assertFalse(TerminacionFibra.objects.exists())

    def test_reimportar_una_reserva_es_idempotente(self):
        reservar_puerto(puerto_id=self.puertos[0].pk)
        fila = self._fila('1', 'Reservar')
        validacion = self._post([fila], 'validar').json()
        self.assertTrue(validacion['can_apply'])
        self.assertEqual(validacion['rows'][0]['resultado'], 'SIN_CAMBIOS')
        aplicacion = self._post([fila], 'aplicar').json()
        self.assertEqual(aplicacion['summary']['sin_cambios'], 1)
        self.odf_a.refresh_from_db()
        self.assertEqual(self.odf_a.puertos_reservados, 1)

    def test_excel_tambien_cancela_reserva_y_desconecta_con_sincronizacion(self):
        reservar_puerto(puerto_id=self.puertos[0].pk)
        conectar_puerto(
            puerto_id=self.puertos[1].pk,
            fibra_id=self.fibra.pk,
            extremo='A',
            ocupar_fibra=True,
        )
        respuesta = self._post([
            self._fila('1', 'Cancelar reserva'),
            self._fila('2', 'Desconectar', sincronizar='Sí'),
        ], 'aplicar')
        self.assertEqual(respuesta.status_code, 200)
        self.puertos[0].refresh_from_db()
        self.puertos[1].refresh_from_db()
        self.fibra.refresh_from_db()
        self.assertEqual(self.puertos[0].estado_puerto, 'LIBRE')
        self.assertEqual(self.puertos[1].estado_puerto, 'LIBRE')
        self.assertEqual(self.fibra.estado, 'DISPONIBLE')
        self.assertEqual(
            set(self.fibra.asignaciones_tramo.values_list('estado', flat=True)),
            {'DISPONIBLE'},
        )
        self.assertFalse(TerminacionFibra.objects.exists())

    def test_error_de_regla_operativa_revierte_la_simulacion_completa(self):
        DetallePuertoODF.objects.filter(pk=self.puertos[1].pk).update(
            estado_puerto='OCUPADO'
        )
        filas = [
            self._fila('1', 'Reservar'),
            self._fila('2', 'Reservar'),
        ]
        validacion = self._post(filas, 'validar')
        self.assertFalse(validacion.json()['can_apply'])
        self.puertos[0].refresh_from_db()
        self.assertEqual(self.puertos[0].estado_puerto, 'LIBRE')

        aplicacion = self._post(filas, 'aplicar')
        self.assertEqual(aplicacion.status_code, 409)
        self.puertos[0].refresh_from_db()
        self.assertEqual(self.puertos[0].estado_puerto, 'LIBRE')

    def test_mover_terminacion_exige_autorizacion_explicita(self):
        TerminacionFibra.objects.create(
            fibra=self.fibra,
            extremo='A',
            puerto_odf=self.puertos[2],
        )
        sin_permiso = self._post([
            self._fila('2', 'Conectar', 'RUTA-EXCEL', 'F12', 'A', 'No', 'No')
        ])
        self.assertFalse(sin_permiso.json()['can_apply'])
        self.assertIn('Permitir mover', sin_permiso.json()['rows'][0]['mensaje'])

        con_permiso = self._post([
            self._fila('2', 'Conectar', 'RUTA-EXCEL', 'F12', 'A', 'No', 'Sí')
        ], 'aplicar')
        self.assertEqual(con_permiso.status_code, 200)
        self.assertEqual(
            TerminacionFibra.objects.get(fibra=self.fibra, extremo='A').puerto_odf_id,
            self.puertos[1].pk,
        )

    def test_no_permite_repetir_puerto_en_el_mismo_archivo(self):
        respuesta = self._post([
            self._fila('1', 'Reservar'),
            self._fila('1', 'Cancelar reserva'),
        ])
        self.assertFalse(respuesta.json()['can_apply'])
        self.assertIn('fila 2', respuesta.json()['rows'][1]['mensaje'])

    def test_valida_permisos_segun_las_acciones_del_archivo(self):
        self.client.force_login(self.usuario_limitado)
        respuesta = self._post([
            self._fila('2', 'Conectar', 'RUTA-EXCEL', 'F12', 'A')
        ])
        self.assertEqual(respuesta.status_code, 403)
        self.assertFalse(TerminacionFibra.objects.exists())

    def test_pantalla_expone_flujo_de_validar_y_aplicar(self):
        respuesta = self.client.get(reverse('planta_interna'))
        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, 'id="ports-bulk-import"')
        self.assertContains(respuesta, 'Validar archivo')
        self.assertContains(respuesta, 'Aplicar archivo')
        self.assertContains(respuesta, reverse('plantilla_operaciones_puertos'))
