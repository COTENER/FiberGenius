"""Actualización masiva de puertos usando las reglas operativas de la GUI."""

from __future__ import annotations

import io
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from openpyxl import load_workbook

from ..models import (
    DetallePuertoODF,
    InventarioFibra,
    InventarioODF,
    Ruta,
    TerminacionFibra,
)
from .puertos import (
    MovimientoRequiereConfirmacion,
    cancelar_reserva_puerto,
    conectar_puerto,
    desconectar_puerto,
    reservar_puerto,
)
from .ubicacion import nombre_site


ACCIONES = {
    "conectar": "Conectar",
    "desconectar": "Desconectar",
    "reservar": "Reservar",
    "cancelar_reserva": "Cancelar reserva",
}

ALIASES_ACCION = {
    "conectar": "conectar",
    "desconectar": "desconectar",
    "reservar": "reservar",
    "cancelar_reserva": "cancelar_reserva",
    "cancelar": "cancelar_reserva",
}

ALIASES_COLUMNAS = {
    "site": "site",
    "sitio": "site",
    "odf": "odf",
    "puerto": "puerto",
    "puerto_odf": "puerto",
    "accion": "accion",
    "troncal": "troncal",
    "ruta": "troncal",
    "fibra": "fibra",
    "hilo": "fibra",
    "codigo_fibra": "codigo_fibra",
    "codigo_de_fibra": "codigo_fibra",
    "extremo": "extremo",
    "sincronizar_fibra": "sincronizar_fibra",
    "sincronizar": "sincronizar_fibra",
    "permitir_mover": "permitir_mover",
}


@dataclass(frozen=True)
class OperacionPuerto:
    fila: int
    accion: str
    site: str
    odf: str
    puerto: str
    puerto_id: int
    troncal: str = ""
    fibra: str = ""
    codigo_fibra: str = ""
    fibra_id: int | None = None
    extremo: str = ""
    sincronizar_fibra: bool = False
    permitir_mover: bool = False

    def como_dict(self):
        return {
            "fila": self.fila,
            "accion": ACCIONES[self.accion],
            "site": self.site,
            "odf": self.odf,
            "puerto": self.puerto,
            "troncal": self.troncal,
            "fibra": self.fibra,
            "codigo_fibra": self.codigo_fibra,
            "extremo": self.extremo,
            "sincronizar_fibra": self.sincronizar_fibra,
            "permitir_mover": self.permitir_mover,
        }


def _normalizar(valor):
    texto = unicodedata.normalize("NFKD", str(valor or "").strip())
    texto = "".join(item for item in texto if not unicodedata.combining(item))
    return re.sub(r"[^a-z0-9]+", "_", texto.casefold()).strip("_")


def _texto(valor):
    if valor is None:
        return ""
    texto = str(valor).strip()
    if texto.casefold() in {"", "nan", "none", "null", "n/a", "-"}:
        return ""
    return texto


def _booleano(valor, columna):
    texto = _normalizar(valor)
    if not texto:
        return False
    if texto in {"si", "s", "true", "verdadero", "1", "yes"}:
        return True
    if texto in {"no", "n", "false", "falso", "0"}:
        return False
    raise ValidationError(f"{columna} debe contener Sí o No.")


def leer_operaciones_excel(contenido, *, max_filas=5000):
    """Lee exclusivamente la hoja Operaciones y devuelve filas normalizadas."""
    try:
        workbook = load_workbook(
            io.BytesIO(contenido),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise ValidationError("El archivo Excel no es válido o está dañado.") from exc
    try:
        if "Operaciones" not in workbook.sheetnames:
            raise ValidationError("El Excel debe contener una hoja llamada Operaciones.")
        sheet = workbook["Operaciones"]
        iterator = sheet.iter_rows(values_only=True)
        try:
            encabezados_originales = next(iterator)
        except StopIteration as exc:
            raise ValidationError("La hoja Operaciones está vacía.") from exc

        encabezados = []
        repetidos = set()
        for original in encabezados_originales:
            clave = ALIASES_COLUMNAS.get(_normalizar(original), _normalizar(original))
            if clave and clave in encabezados:
                repetidos.add(clave)
            encabezados.append(clave)
        if repetidos:
            raise ValidationError(
                "Hay columnas equivalentes repetidas: " + ", ".join(sorted(repetidos)) + "."
            )
        obligatorias = {"odf", "puerto", "accion"}
        faltantes = obligatorias - set(encabezados)
        if faltantes:
            raise ValidationError(
                "Faltan columnas obligatorias: " + ", ".join(sorted(faltantes)) + "."
            )

        filas = []
        for numero_fila, valores in enumerate(iterator, start=2):
            registro = {
                encabezados[indice]: valor
                for indice, valor in enumerate(valores)
                if indice < len(encabezados) and encabezados[indice]
            }
            if not any(_texto(valor) for valor in registro.values()):
                continue
            filas.append((numero_fila, registro))
            if len(filas) > max_filas:
                raise ValidationError(
                    f"El archivo supera el máximo permitido de {max_filas} operaciones."
                )
        if not filas:
            raise ValidationError("La hoja Operaciones no contiene filas para procesar.")
        return filas
    finally:
        workbook.close()


def _error_fila(fila, mensaje, registro=None):
    registro = registro or {}
    return {
        "fila": fila,
        "resultado": "ERROR",
        "accion": _texto(registro.get("accion")),
        "site": _texto(registro.get("site")),
        "odf": _texto(registro.get("odf")),
        "puerto": _texto(registro.get("puerto")),
        "mensaje": mensaje,
    }


def preparar_operaciones(filas):
    """Resuelve identidades y rechaza ambigüedades antes de cambiar estados."""
    operaciones = []
    resultados = []
    puertos_vistos = {}
    extremos_vistos = {}

    # La resolución es masiva: el costo de consultas no debe crecer por cada
    # fila del Excel. También mantiene una fotografía coherente durante toda
    # la validación previa.
    odfs_por_nombre = defaultdict(list)
    for odf_obj in InventarioODF.objects.select_related(
        "rack_obj__sala__hub_site"
    ).order_by("pk"):
        odfs_por_nombre[odf_obj.odf.casefold()].append(odf_obj)
    odf_ids = {
        odf_obj.pk
        for _, registro in filas
        for odf_obj in odfs_por_nombre.get(
            _texto(registro.get("odf")).casefold(),
            (),
        )
    }
    puertos_por_identidad = defaultdict(list)
    for puerto_obj in DetallePuertoODF.objects.filter(
        odf_obj_id__in=odf_ids
    ).select_related("odf_obj"):
        puertos_por_identidad[
            (puerto_obj.odf_obj_id, puerto_obj.puerto_odf.casefold())
        ].append(puerto_obj)
    rutas_por_nombre = {
        ruta.nombre.casefold(): ruta for ruta in Ruta.objects.all()
    }
    fibras_por_codigo = defaultdict(list)
    fibras_por_ruta_numero = defaultdict(list)
    for fibra_obj in InventarioFibra.objects.select_related("ruta"):
        fibras_por_codigo[fibra_obj.codigo_fibra.casefold()].append(fibra_obj)
        if fibra_obj.ruta_id:
            fibras_por_ruta_numero[
                (fibra_obj.ruta_id, fibra_obj.fibra_numero.casefold())
            ].append(fibra_obj)

    for numero_fila, registro in filas:
        try:
            site = _texto(registro.get("site"))
            odf_nombre = _texto(registro.get("odf"))
            puerto_nombre = _texto(registro.get("puerto"))
            accion = ALIASES_ACCION.get(_normalizar(registro.get("accion")))
            if not odf_nombre or not puerto_nombre:
                raise ValidationError("ODF y Puerto son obligatorios.")
            if not accion:
                raise ValidationError(
                    "Acción no válida. Use Conectar, Desconectar, Reservar o Cancelar reserva."
                )

            odfs = odfs_por_nombre.get(odf_nombre.casefold(), [])
            if not odfs:
                raise ValidationError(f"No existe el ODF {odf_nombre}.")
            if len(odfs) > 1:
                raise ValidationError(f"El ODF {odf_nombre} es ambiguo.")
            site_real = nombre_site(odfs[0])
            if site and site.casefold() != site_real.casefold():
                raise ValidationError(
                    f"El Site {site} no corresponde al ODF {odf_nombre}; "
                    f"pertenece a {site_real}."
                )
            puertos = puertos_por_identidad.get(
                (odfs[0].pk, puerto_nombre.casefold()),
                [],
            )
            if not puertos:
                raise ValidationError(
                    f"No existe el puerto {puerto_nombre} en {site} / {odf_nombre}."
                )
            if len(puertos) > 1:
                raise ValidationError(
                    f"El puerto {puerto_nombre} es ambiguo en el ODF {odf_nombre}."
                )
            puerto = puertos[0]
            if puerto.pk in puertos_vistos:
                raise ValidationError(
                    f"El puerto ya fue incluido en la fila {puertos_vistos[puerto.pk]}."
                )

            troncal = _texto(registro.get("troncal"))
            fibra_numero = _texto(registro.get("fibra")).upper()
            codigo_fibra = _texto(registro.get("codigo_fibra")).upper()
            extremo = _texto(registro.get("extremo")).upper()
            fibra = None
            if accion == "conectar":
                if extremo not in {"A", "B"}:
                    raise ValidationError(
                        "Conectar requiere Extremo A o B."
                    )
                if codigo_fibra:
                    candidatas_codigo = fibras_por_codigo.get(
                        codigo_fibra.casefold(),
                        [],
                    )
                    if not candidatas_codigo:
                        raise ValidationError(
                            f"No existe el Codigo Fibra {codigo_fibra}."
                        )
                    if len(candidatas_codigo) > 1:
                        raise ValidationError(
                            f"El Codigo Fibra {codigo_fibra} es ambiguo."
                        )
                    fibra = candidatas_codigo[0]
                    if (
                        troncal
                        and fibra.ruta_id
                        and fibra.ruta.nombre.casefold() != troncal.casefold()
                    ):
                        raise ValidationError(
                            f"El Codigo Fibra {codigo_fibra} pertenece a "
                            f"{fibra.ruta.nombre}, no a {troncal}."
                        )
                    if (
                        fibra_numero
                        and fibra.fibra_numero.casefold()
                        != fibra_numero.casefold()
                    ):
                        raise ValidationError(
                            f"El Codigo Fibra {codigo_fibra} corresponde a "
                            f"{fibra.fibra_numero}, no a {fibra_numero}."
                        )
                else:
                    if not troncal or not fibra_numero:
                        raise ValidationError(
                            "Conectar requiere Codigo Fibra. Troncal + Fibra "
                            "se acepta solo por compatibilidad legacy."
                        )
                    ruta = rutas_por_nombre.get(troncal.casefold())
                    if ruta is None:
                        raise ValidationError(f"No existe la troncal {troncal}.")
                    candidatas = fibras_por_ruta_numero.get(
                        (ruta.pk, fibra_numero.casefold()),
                        [],
                    )
                    if not candidatas:
                        raise ValidationError(
                            f"No existe la fibra {fibra_numero} en la troncal {troncal}."
                        )
                    if len(candidatas) > 1:
                        raise ValidationError(
                            f"{troncal} / {fibra_numero} es ambiguo; informe Codigo Fibra."
                        )
                    fibra = candidatas[0]
                    codigo_fibra = fibra.codigo_fibra
                clave_extremo = (fibra.pk, extremo)
                if clave_extremo in extremos_vistos:
                    raise ValidationError(
                        f"Ese extremo ya fue incluido en la fila {extremos_vistos[clave_extremo]}."
                    )
                extremos_vistos[clave_extremo] = numero_fila

            sincronizar_fibra = _booleano(
                registro.get("sincronizar_fibra"),
                "Sincronizar fibra",
            )
            permitir_mover = _booleano(
                registro.get("permitir_mover"),
                "Permitir mover",
            )
            if sincronizar_fibra and accion not in {"conectar", "desconectar"}:
                raise ValidationError(
                    "Sincronizar fibra solo aplica a Conectar o Desconectar."
                )
            if permitir_mover and accion != "conectar":
                raise ValidationError("Permitir mover solo aplica a Conectar.")

            operacion = OperacionPuerto(
                fila=numero_fila,
                accion=accion,
                site=nombre_site(odfs[0]),
                odf=odfs[0].odf,
                puerto=puerto.puerto_odf,
                puerto_id=puerto.pk,
                troncal=troncal,
                fibra=fibra_numero,
                codigo_fibra=codigo_fibra,
                fibra_id=fibra.pk if fibra else None,
                extremo=extremo,
                sincronizar_fibra=sincronizar_fibra,
                permitir_mover=permitir_mover,
            )
            puertos_vistos[puerto.pk] = numero_fila
            operaciones.append(operacion)
        except ValidationError as exc:
            mensaje = "; ".join(exc.messages)
            resultados.append(_error_fila(numero_fila, mensaje, registro))
    return operaciones, resultados


def permisos_requeridos(operaciones):
    permisos = {"mapas.change_detallepuertoodf"}
    for operacion in operaciones:
        if operacion.accion == "conectar":
            # Un mismo archivo puede desconectar y volver a conectar un extremo;
            # por ello la carga masiva requiere ambas capacidades.
            permisos.update({
                "mapas.add_terminacionfibra",
                "mapas.change_terminacionfibra",
            })
        elif operacion.accion == "desconectar":
            permisos.add("mapas.delete_terminacionfibra")
        if operacion.sincronizar_fibra:
            permisos.add("mapas.change_inventariofibra")
    return permisos


def _fibra_ya_en_estado(fibra_id, estado):
    return InventarioFibra.objects.filter(
        pk=fibra_id,
        estado=estado,
        origen_estado="INFORMADO",
    ).exists()


def _ejecutar_operacion(
    operacion, *, usuario=None, origen="EXCEL", lote_importacion=None
):
    puerto = DetallePuertoODF.objects.select_related("odf_obj").get(
        pk=operacion.puerto_id
    )
    terminacion_puerto = TerminacionFibra.objects.filter(
        puerto_odf_id=puerto.pk
    ).first()

    if operacion.accion == "reservar":
        if puerto.estado_puerto == "RESERVADO" and terminacion_puerto is None:
            return "SIN_CAMBIOS", "El puerto ya estaba reservado."
        reservar_puerto(
            puerto_id=puerto.pk,
            usuario=usuario,
            origen=origen,
            lote_importacion=lote_importacion,
        )
        return "APLICADA", "Puerto reservado."

    if operacion.accion == "cancelar_reserva":
        if puerto.estado_puerto == "LIBRE" and terminacion_puerto is None:
            return "SIN_CAMBIOS", "El puerto ya estaba libre."
        cancelar_reserva_puerto(
            puerto_id=puerto.pk,
            usuario=usuario,
            origen=origen,
            lote_importacion=lote_importacion,
        )
        return "APLICADA", "Reserva cancelada."

    if operacion.accion == "desconectar":
        if puerto.estado_puerto == "LIBRE" and terminacion_puerto is None:
            return "SIN_CAMBIOS", "El puerto ya estaba libre y sin terminación."
        desconectar_puerto(
            puerto_id=puerto.pk,
            liberar_fibra=operacion.sincronizar_fibra,
            usuario=usuario,
            origen=origen,
            lote_importacion=lote_importacion,
        )
        mensaje = "Puerto desconectado."
        if operacion.sincronizar_fibra:
            mensaje += " Estado global de la fibra informado como Disponible."
        return "APLICADA", mensaje

    existente = TerminacionFibra.objects.filter(
        fibra_id=operacion.fibra_id,
        extremo=operacion.extremo,
    ).first()
    misma_conexion = bool(
        existente and existente.puerto_odf_id == puerto.pk
        and puerto.estado_puerto == "OCUPADO"
    )
    sincronizacion_pendiente = bool(
        operacion.sincronizar_fibra
        and not _fibra_ya_en_estado(operacion.fibra_id, "OCUPADO")
    )
    if misma_conexion and not sincronizacion_pendiente:
        return "SIN_CAMBIOS", "La conexión ya estaba registrada."
    conectar_puerto(
        puerto_id=puerto.pk,
        fibra_id=operacion.fibra_id,
        extremo=operacion.extremo,
        permitir_mover=operacion.permitir_mover,
        ocupar_fibra=operacion.sincronizar_fibra,
        usuario=usuario,
        origen=origen,
        lote_importacion=lote_importacion,
    )
    mensaje = "Conexión registrada."
    if operacion.sincronizar_fibra:
        mensaje += " Estado global de la fibra informado como Ocupado."
    return "APLICADA", mensaje


def ejecutar_operaciones(
    operaciones, *, aplicar=False, usuario=None, origen="EXCEL",
    lote_importacion=None,
):
    """Ejecuta los servicios reales y revierte siempre cuando solo valida."""
    resultados = []
    errores = 0
    with transaction.atomic():
        for operacion in operaciones:
            base = operacion.como_dict()
            try:
                resultado, mensaje = _ejecutar_operacion(
                    operacion,
                    usuario=usuario,
                    origen=origen,
                    lote_importacion=lote_importacion,
                )
                resultados.append({
                    **base,
                    "resultado": (
                        resultado if aplicar or resultado == "SIN_CAMBIOS" else "VALIDA"
                    ),
                    "mensaje": (
                        mensaje if aplicar or resultado == "SIN_CAMBIOS"
                        else "Operación válida."
                    ),
                })
            except MovimientoRequiereConfirmacion as exc:
                errores += 1
                resultados.append({
                    **base,
                    "resultado": "ERROR",
                    "mensaje": "; ".join(exc.messages)
                    + " Use Permitir mover = Sí para autorizarlo.",
                })
            except (ValidationError, IntegrityError) as exc:
                errores += 1
                mensajes = exc.messages if isinstance(exc, ValidationError) else [str(exc)]
                resultados.append({
                    **base,
                    "resultado": "ERROR",
                    "mensaje": "; ".join(mensajes),
                })
        if not aplicar or errores:
            transaction.set_rollback(True)
    return resultados


def resumen_resultados(resultados):
    return {
        "total": len(resultados),
        "validas": sum(item["resultado"] == "VALIDA" for item in resultados),
        "aplicadas": sum(item["resultado"] == "APLICADA" for item in resultados),
        "sin_cambios": sum(item["resultado"] == "SIN_CAMBIOS" for item in resultados),
        "no_aplicadas": sum(item["resultado"] == "NO_APLICADA" for item in resultados),
        "errores": sum(item["resultado"] == "ERROR" for item in resultados),
    }
