"""Prevalidación y aplicación atómica de operaciones ODF por Excel."""

from __future__ import annotations

import hashlib
import io
import logging
import os
import zipfile

from django.conf import settings
from django.contrib.auth.decorators import login_required, permission_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404
from django.urls import reverse
from django.utils.timezone import now
from django.views.decorators.http import require_GET, require_POST
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from ..models import LoteImportacion
from ..services.importacion_puertos import (
    ejecutar_operaciones,
    leer_operaciones_excel,
    permisos_requeridos,
    preparar_operaciones,
    resumen_resultados,
)


logger = logging.getLogger("fibergenius.import")
TIPO_LOTE = "operaciones_puertos_excel"
CONTENT_TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _respuesta_excel(workbook, nombre):
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=CONTENT_TYPE_XLSX)
    response["Content-Disposition"] = f'attachment; filename="{nombre}.xlsx"'
    return response


def _estilo_encabezado(celdas):
    borde = Border(bottom=Side(style="thin", color="A7B4C8"))
    for celda in celdas:
        celda.fill = PatternFill("solid", fgColor="173B68")
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde


@login_required
@permission_required("mapas.change_detallepuertoodf", raise_exception=True)
@require_GET
def descargar_plantilla_operaciones_puertos(request):
    workbook = Workbook()
    operaciones = workbook.active
    operaciones.title = "Operaciones"
    encabezados = [
        "ODF",
        "Puerto",
        "Acción",
        "Codigo Fibra",
        "Extremo",
        "Sincronizar fibra",
        "Permitir mover",
    ]
    operaciones.append(encabezados)
    _estilo_encabezado(operaciones[1])
    operaciones.freeze_panes = "A2"
    operaciones.auto_filter.ref = "A1:G1"
    operaciones.sheet_view.showGridLines = False
    operaciones.row_dimensions[1].height = 28
    anchos = (28, 13, 22, 36, 13, 22, 18)
    for indice, ancho in enumerate(anchos, start=1):
        operaciones.column_dimensions[chr(64 + indice)].width = ancho

    ayudas = {
        "C1": "Conectar, Desconectar, Reservar o Cancelar reserva.",
        "D1": "Obligatorio solamente para Conectar; identifica la fibra sin ambigüedad.",
        "E1": "A o B. Obligatorio solamente para Conectar.",
        "F1": "Sí informa también el estado global de InventarioFibra.",
        "G1": "Sí autoriza mover un extremo ya conectado a otro puerto.",
    }
    for referencia, texto in ayudas.items():
        operaciones[referencia].comment = Comment(texto, "Fiber Genius")

    validacion_accion = DataValidation(
        type="list",
        formula1='"Conectar,Desconectar,Reservar,Cancelar reserva"',
        allow_blank=False,
    )
    validacion_extremo = DataValidation(type="list", formula1='"A,B"', allow_blank=True)
    validacion_si_no = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    operaciones.add_data_validation(validacion_accion)
    operaciones.add_data_validation(validacion_extremo)
    operaciones.add_data_validation(validacion_si_no)
    validacion_accion.add("C2:C5001")
    validacion_extremo.add("E2:E5001")
    validacion_si_no.add("F2:G5001")

    instrucciones = workbook.create_sheet("Instrucciones")
    instrucciones.sheet_view.showGridLines = False
    instrucciones.merge_cells("A1:F1")
    instrucciones["A1"] = "Fiber Genius · Actualización masiva de puertos ODF"
    instrucciones["A1"].fill = PatternFill("solid", fgColor="0F2747")
    instrucciones["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    instrucciones["A1"].alignment = Alignment(vertical="center")
    instrucciones.row_dimensions[1].height = 34
    instrucciones["A3"] = "Flujo seguro"
    instrucciones["A3"].font = Font(bold=True, color="173B68", size=12)
    pasos = [
        "1. Complete solamente la hoja Operaciones.",
        "2. Descargue o conserve esta plantilla en formato .xlsx.",
        "3. En Fiber Genius seleccione Validar archivo.",
        "4. Revise el resultado por fila y aplique solo cuando no existan errores.",
        "5. La aplicación es atómica: si una fila falla, no se guarda ninguna.",
    ]
    for fila, paso in enumerate(pasos, start=4):
        instrucciones.cell(fila, 1, paso)
    instrucciones["A11"] = "Reglas"
    instrucciones["A11"].font = Font(bold=True, color="173B68", size=12)
    reglas = [
        "ODF + Puerto identifican el puerto físico.",
        "Conectar exige Codigo Fibra y Extremo.",
        "Sincronizar fibra y Permitir mover asumen No cuando están vacíos.",
        "El archivo no puede repetir el mismo puerto ni el mismo extremo de una fibra.",
        "Las operaciones reutilizan las mismas reglas transaccionales de la GUI.",
    ]
    for fila, regla in enumerate(reglas, start=12):
        instrucciones.cell(fila, 1, regla)
    instrucciones.column_dimensions["A"].width = 95
    instrucciones.column_dimensions["B"].width = 4

    ejemplo = workbook.create_sheet("Ejemplo - no importar")
    ejemplo.append(encabezados)
    _estilo_encabezado(ejemplo[1])
    ejemplo.append(["ODF-01", "1", "Desconectar", "", "", "No", "No"])
    ejemplo.append(["ODF-01", "2", "Reservar", "", "", "No", "No"])
    ejemplo.append(["ODF-01", "3", "Conectar", "FGF-000012", "A", "Sí", "No"])
    ejemplo.freeze_panes = "A2"
    ejemplo.sheet_view.showGridLines = False
    for indice, ancho in enumerate(anchos, start=1):
        ejemplo.column_dimensions[chr(64 + indice)].width = ancho

    workbook.active = 0
    return _respuesta_excel(workbook, "FiberGenius_Operaciones_Puertos")


def _validar_contenido_xlsx(archivo):
    if not archivo.name.casefold().endswith(".xlsx"):
        raise ValidationError("El archivo debe tener formato .xlsx.")
    limite = int(getattr(settings, "FIBERGENIUS_MAX_UPLOAD_BYTES", 25 * 1024 * 1024))
    if archivo.size > limite:
        raise ValidationError(
            f"El archivo supera el límite permitido de {limite // (1024 * 1024)} MB."
        )
    contenido = archivo.read()
    try:
        with zipfile.ZipFile(io.BytesIO(contenido), "r") as paquete:
            entradas = [item for item in paquete.infolist() if not item.is_dir()]
            max_entradas = int(getattr(settings, "FIBERGENIUS_MAX_ZIP_ENTRIES", 500))
            max_total = int(
                getattr(settings, "FIBERGENIUS_MAX_ZIP_UNCOMPRESSED_BYTES", 100 * 1024 * 1024)
            )
            if len(entradas) > max_entradas:
                raise ValidationError(f"El Excel contiene más de {max_entradas} entradas internas.")
            if sum(item.file_size for item in entradas) > max_total:
                raise ValidationError("El contenido descomprimido del Excel supera el límite permitido.")
    except zipfile.BadZipFile as exc:
        raise ValidationError("El archivo .xlsx no es válido.") from exc
    return contenido


def _combinar_resultados(errores_preparacion, resultados_ejecucion):
    return sorted(
        [*errores_preparacion, *resultados_ejecucion],
        key=lambda item: item["fila"],
    )


@login_required
@permission_required("mapas.change_detallepuertoodf", raise_exception=True)
@require_POST
def importar_operaciones_puertos(request):
    archivo = request.FILES.get("archivo")
    modo = str(request.POST.get("modo", "validar")).strip().casefold()
    if archivo is None:
        return JsonResponse(
            {"status": "error", "message": "Seleccione un archivo Excel."},
            status=400,
        )
    if modo not in {"validar", "aplicar"}:
        return JsonResponse(
            {"status": "error", "message": "El modo solicitado no es válido."},
            status=400,
        )

    try:
        contenido = _validar_contenido_xlsx(archivo)
        filas = leer_operaciones_excel(
            contenido,
            max_filas=int(getattr(settings, "FIBERGENIUS_MAX_BATCH_PORT_OPERATIONS", 5000)),
        )
        operaciones, errores_preparacion = preparar_operaciones(filas)
        permisos = permisos_requeridos(operaciones)
        if not request.user.is_superuser and not request.user.has_perms(permisos):
            raise PermissionDenied

        simulados = ejecutar_operaciones(
            operaciones,
            aplicar=False,
            usuario=request.user,
            origen="EXCEL",
        )
        resultados_validacion = _combinar_resultados(errores_preparacion, simulados)
        resumen_validacion = resumen_resultados(resultados_validacion)
        puede_aplicar = resumen_validacion["errores"] == 0
        if modo == "validar" or not puede_aplicar:
            return JsonResponse(
                {
                    "status": "valid" if puede_aplicar else "invalid",
                    "message": (
                        "Todas las operaciones son válidas. Ya puede aplicar el archivo."
                        if puede_aplicar
                        else "El archivo contiene errores; no se realizó ningún cambio."
                    ),
                    "can_apply": puede_aplicar,
                    "summary": resumen_validacion,
                    "rows": resultados_validacion,
                },
                status=200 if modo == "validar" else 409,
            )

        lote = LoteImportacion.objects.create(
            tipo=TIPO_LOTE,
            archivo_origen=os.path.basename(archivo.name),
            hash_sha256=hashlib.sha256(contenido).hexdigest(),
            estado="PROCESANDO",
            usuario=request.user,
            total_filas=len(resultados_validacion),
        )
        try:
            aplicados = ejecutar_operaciones(
                operaciones,
                aplicar=True,
                usuario=request.user,
                origen="EXCEL",
                lote_importacion=lote,
            )
            resultados = _combinar_resultados([], aplicados)
            resumen = resumen_resultados(resultados)
            if resumen["errores"]:
                for item in resultados:
                    if item["resultado"] != "ERROR":
                        item["resultado"] = "NO_APLICADA"
                        item["mensaje"] = "No aplicada porque otra fila del archivo falló."
                resumen = resumen_resultados(resultados)
                lote.estado = "FALLIDO"
            else:
                lote.estado = "COMPLETADO"
            lote.filas_actualizadas = resumen["aplicadas"]
            lote.filas_rechazadas = resumen["errores"]
            lote.detalle_errores = [
                {"fila": item["fila"], "mensaje": item["mensaje"]}
                for item in resultados if item["resultado"] == "ERROR"
            ]
            lote.metadatos_origen = {
                "resumen": resumen,
                "resultados": resultados,
                "flujo": "servicios_operativos_puertos",
            }
            lote.finalizado_en = now()
            lote.save(update_fields=[
                "estado",
                "filas_actualizadas",
                "filas_rechazadas",
                "detalle_errores",
                "metadatos_origen",
                "finalizado_en",
            ])
            correcto = lote.estado == "COMPLETADO"
            return JsonResponse(
                {
                    "status": "success" if correcto else "error",
                    "message": (
                        "Archivo aplicado completamente."
                        if correcto else "El archivo no fue aplicado porque una operación falló."
                    ),
                    "can_apply": False,
                    "summary": resumen,
                    "rows": resultados,
                    "lote": str(lote.codigo),
                    "report_url": reverse(
                        "resultado_operaciones_puertos",
                        args=[lote.codigo],
                    ),
                },
                status=200 if correcto else 409,
            )
        except Exception:
            logger.exception("Falló la aplicación masiva de puertos")
            lote.estado = "FALLIDO"
            lote.detalle_errores = [{"mensaje": "Error inesperado durante la aplicación."}]
            lote.finalizado_en = now()
            lote.save(update_fields=["estado", "detalle_errores", "finalizado_en"])
            return JsonResponse(
                {
                    "status": "error",
                    "message": "No se aplicó ninguna operación. Consulte el lote generado.",
                },
                status=500,
            )
    except PermissionDenied:
        raise
    except ValidationError as exc:
        return JsonResponse(
            {"status": "error", "message": "; ".join(exc.messages)},
            status=400,
        )
    except Exception:
        logger.exception("No se pudo procesar el Excel de operaciones ODF")
        return JsonResponse(
            {"status": "error", "message": "No se pudo procesar el archivo Excel."},
            status=500,
        )


@login_required
@permission_required("mapas.change_detallepuertoodf", raise_exception=True)
@require_GET
def descargar_resultado_operaciones_puertos(request, codigo):
    lote = get_object_or_404(LoteImportacion, codigo=codigo, tipo=TIPO_LOTE)
    if not request.user.is_superuser and lote.usuario_id != request.user.pk:
        raise PermissionDenied
    resultados = list((lote.metadatos_origen or {}).get("resultados") or [])
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Resultado"
    encabezados = [
        "Fila", "Resultado", "Acción", "Site", "ODF", "Puerto",
        "Troncal", "Fibra", "Codigo Fibra", "Extremo", "Mensaje",
    ]
    sheet.append(encabezados)
    _estilo_encabezado(sheet[1])

    def valor_excel_seguro(valor):
        # Evita que texto procedente del archivo se interprete como una fórmula
        # al abrir el reporte en Excel.
        if isinstance(valor, str) and valor.startswith(("=", "+", "-", "@")):
            return "'" + valor
        return valor

    for item in resultados:
        sheet.append([valor_excel_seguro(valor) for valor in [
            item.get("fila"),
            item.get("resultado"),
            item.get("accion"),
            item.get("site"),
            item.get("odf"),
            item.get("puerto"),
            item.get("troncal"),
            item.get("fibra"),
            item.get("codigo_fibra"),
            item.get("extremo"),
            item.get("mensaje"),
        ]])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:K{max(sheet.max_row, 1)}"
    sheet.sheet_view.showGridLines = False
    for columna, ancho in zip("ABCDEFGHIJK", (9, 16, 20, 24, 28, 13, 30, 14, 36, 12, 70)):
        sheet.column_dimensions[columna].width = ancho
    return _respuesta_excel(workbook, f"FiberGenius_Resultado_{lote.codigo}")
